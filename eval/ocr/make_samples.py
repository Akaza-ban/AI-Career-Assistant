"""
Tạo bộ CV mẫu hư cấu cho bộ đo đọc CV (eval/ocr/samples):
- <mẫu>.pdf: PDF có lớp chữ (render_samples.cjs dựng từ HTML)
- <mẫu>.scan.pdf: chỉ có ảnh trang ~150 dpi, như máy scan
- <mẫu>.photo.pdf: ảnh nghiêng nhẹ, hơi mờ, nén JPEG, độ phân giải thấp, như chụp bằng điện thoại
- vi_docx.docx (đoạn văn, tab ngày, bảng) và en_textbox.docx (thông tin nằm trong hộp chữ)
- truth.json: đáp án (chữ hiển thị, họ tên, tiêu đề mục) cho từng mẫu; files.json: danh sách file cần đo.

Chạy từ thư mục gốc:  python -m eval.ocr.make_samples   (cần node + playwright trong frontend/node_modules)
"""

import io
import json
import random
import subprocess
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.shared import Pt
from PIL import Image, ImageEnhance, ImageFilter

HERE = Path(__file__).parent
SAMPLES = HERE / "samples"


def _image_pdf(img: Image.Image, *, jpeg_quality: int | None = None) -> bytes:
    if jpeg_quality:
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=jpeg_quality)
        img = Image.open(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    img.convert("RGB").save(out, format="PDF", resolution=150)
    return out.getvalue()


def scan(png: Path) -> bytes:
    """Ảnh trang A4 ~150 dpi, thang xám nhẹ như máy scan văn phòng."""
    img = Image.open(png).convert("L").resize((1240, 1754), Image.LANCZOS)
    return _image_pdf(img.convert("RGB"))


def photo(png: Path, seed: int) -> bytes:
    """Như chụp điện thoại: nghiêng 1–2 độ, mờ nhẹ, tối và nhạt hơn, ~110 dpi, nén JPEG."""
    rnd = random.Random(seed)
    img = Image.open(png).convert("RGB").resize((920, 1300), Image.LANCZOS)
    img = img.rotate(rnd.uniform(-1.8, 1.8), resample=Image.BICUBIC, expand=True, fillcolor=(236, 234, 228))
    img = img.filter(ImageFilter.GaussianBlur(0.7))
    img = ImageEnhance.Brightness(img).enhance(0.93)
    img = ImageEnhance.Contrast(img).enhance(0.85)
    return _image_pdf(img, jpeg_quality=55)


# ---------- DOCX ----------

VI_DOCX = {
    "name": "Đặng Khánh Linh",
    "contact": "Hà Nội | linh.dang@example.com | 0977 456 123",
    "sections": [
        ("MỤC TIÊU", ["Sinh viên năm ba ngành Hệ thống thông tin, mong muốn thực tập ở vị trí Business Analyst."]),
        (
            "HỌC VẤN",
            ["Học viện Công nghệ Bưu chính Viễn thông\t09/2022 – nay", "Hệ thống thông tin quản lý, GPA 3.4/4.0"],
        ),
        (
            "KINH NGHIỆM",
            [
                "Công ty Phần mềm Hòa Bình | Thực tập sinh BA\t06/2024 – 09/2024",
                "• Phỏng vấn 8 người dùng nội bộ để thu thập yêu cầu cho module quản lý kho.",
                "• Viết tài liệu đặc tả yêu cầu cho 12 tính năng và vẽ sơ đồ BPMN.",
            ],
        ),
        ("KỸ NĂNG", []),
    ],
    "table": [("Phân tích nghiệp vụ", "BRD, SRS, BPMN, user story"), ("Công cụ", "Jira, Draw.io, Excel, SQL")],
}

EN_TEXTBOX = {
    "name": "Nguyen Bao Ngoc",
    "boxes": [
        ["CONTACT", "ngoc.nguyen@example.com", "0938 111 222", "Da Nang, Vietnam"],
        ["SKILLS", "Figma, Adobe XD", "User research", "HTML and CSS"],
    ],
    "body": [
        ("PROFILE", ["UI/UX design student who enjoys turning user research into simple, usable interfaces."]),
        (
            "EXPERIENCE",
            [
                "Lotus Studio | UI/UX Design Intern\tJan 2024 – Jun 2024",
                "• Redesigned the booking flow of a travel app, reducing checkout steps from 6 to 3.",
                "• Ran 5 usability tests and presented findings to the product team.",
            ],
        ),
        ("EDUCATION", ["University of Architecture Ho Chi Minh City\t2021 – 2025", "Bachelor of Graphic Design"]),
    ],
}

_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
_V = 'xmlns:v="urn:schemas-microsoft-com:vml"'


def _textbox_run(lines: list[str], top_pt: int) -> object:
    """Một hộp chữ (VML) nổi bên phải trang: nội dung nằm trong w:txbxContent, không phải đoạn văn thường."""
    paras = "".join(f'<w:p><w:r><w:t xml:space="preserve">{t}</w:t></w:r></w:p>' for t in lines)
    return parse_xml(
        f'<w:r {_W}><w:pict><v:shape {_V} id="tb{top_pt}" type="#_x0000_t202" '
        f'style="position:absolute;margin-left:330pt;margin-top:{top_pt}pt;width:160pt;height:110pt;z-index:1">'
        f"<v:textbox><w:txbxContent>{paras}</w:txbxContent></v:textbox></v:shape></w:pict></w:r>"
    )


def make_docx() -> dict[str, dict]:
    truth: dict[str, dict] = {}

    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_paragraph().add_run(VI_DOCX["name"]).bold = True
    doc.add_paragraph(VI_DOCX["contact"])
    lines = [VI_DOCX["name"], VI_DOCX["contact"]]
    for title, rows in VI_DOCX["sections"]:
        doc.add_paragraph().add_run(title).bold = True
        lines.append(title)
        for row in rows:
            doc.add_paragraph(row)
            lines.append(row.replace("\t", " "))
    table = doc.add_table(rows=0, cols=2)
    for label, value in VI_DOCX["table"]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value
        lines.append(f"{label} {value}")
    doc.save(SAMPLES / "vi_docx.docx")
    truth["vi_docx"] = {
        "text": "\n".join(lines),
        "name": VI_DOCX["name"],
        "headings": [t for t, _ in VI_DOCX["sections"]],
    }

    doc = Document()
    doc.add_paragraph().add_run(EN_TEXTBOX["name"]).bold = True
    anchor = doc.add_paragraph()
    for i, box in enumerate(EN_TEXTBOX["boxes"]):
        anchor._p.append(_textbox_run(box, 10 + i * 130))
    lines = [EN_TEXTBOX["name"]] + [t for box in EN_TEXTBOX["boxes"] for t in box]
    for title, rows in EN_TEXTBOX["body"]:
        doc.add_paragraph().add_run(title).bold = True
        lines.append(title)
        for row in rows:
            doc.add_paragraph(row)
            lines.append(row.replace("\t", " "))
    doc.save(SAMPLES / "en_textbox.docx")
    truth["en_textbox"] = {
        "text": "\n".join(lines),
        "name": EN_TEXTBOX["name"],
        "headings": ["SKILLS"] + [t for t, _ in EN_TEXTBOX["body"]],
    }
    return truth


def main() -> None:
    SAMPLES.mkdir(exist_ok=True)
    subprocess.run(["node", str(HERE / "render_samples.cjs"), str(SAMPLES)], check=True)
    truth = json.loads((SAMPLES / "truth.json").read_text(encoding="utf-8"))
    files = []
    for i, name in enumerate(sorted(truth)):
        png = SAMPLES / f"{name}.png"
        (SAMPLES / f"{name}.scan.pdf").write_bytes(scan(png))
        (SAMPLES / f"{name}.photo.pdf").write_bytes(photo(png, seed=i))
        png.unlink()
        files += [
            {"file": f"{name}.pdf", "sample": name, "kind": "pdf"},
            {"file": f"{name}.scan.pdf", "sample": name, "kind": "scan"},
            {"file": f"{name}.photo.pdf", "sample": name, "kind": "photo"},
        ]
    truth.update(make_docx())
    files += [
        {"file": "vi_docx.docx", "sample": "vi_docx", "kind": "docx"},
        {"file": "en_textbox.docx", "sample": "en_textbox", "kind": "docx"},
    ]
    (SAMPLES / "truth.json").write_text(json.dumps(truth, ensure_ascii=False, indent=2), encoding="utf-8")
    (SAMPLES / "files.json").write_text(json.dumps(files, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"ok: {len(files)} file trong {SAMPLES}")


if __name__ == "__main__":
    main()
