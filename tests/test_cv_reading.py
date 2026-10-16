"""Đọc CV: chèn dấu cách OCR bỏ sót, sửa tên công cụ đọc nhầm, tiêu đề sai một ký tự, thứ tự đọc nhiều cột,
DOCX có hộp chữ, PDF hỏng bảng mã font. Chất lượng trên cả bộ CV mẫu đo bằng eval/ocr (python -m eval.ocr.run_ocr_eval)."""

import io

from docx import Document
from docx.oxml import parse_xml

from src.services.cv_files import _extract_docx, looks_garbled, readable_chars
from src.services.cv_layout import extract_layout_aware_text
from src.services.ocr import render_scale, respace
from src.services.ocr_terms import fix_terms, rejoin_headings, static_vocabulary
from src.services.skill_matching import _match_section


def _chars(text: str, gaps: dict[int, float], width: float = 10.0) -> tuple[list[str], list[list[list[float]]]]:
    """Ký tự OCR kèm khung: mỗi ký tự rộng `width`, sát nhau trừ các vị trí trong `gaps` (khoảng trống trước ký tự i)."""
    boxes, x = [], 0.0
    for i, _ in enumerate(text):
        x += gaps.get(i, 1.0)
        boxes.append([[x, 0], [x + width, 0], [x + width, 20], [x, 20]])
        x += width
    return list(text), boxes


def test_respace_inserts_missing_spaces_from_character_gaps():
    text = "abackendengineeringintern"
    # khoảng trống thật trước "backend", "engineering", "intern"
    chars, boxes = _chars(text, {1: 8.0, 8: 8.0, 19: 8.0})
    assert respace(chars, boxes) == "a backend engineering intern"


def test_respace_keeps_emails_and_ignores_noise_on_lines_with_spaces():
    chars, boxes = _chars("mail@example.com", {4: 9.0, 5: 9.0})
    assert respace(chars, boxes) == "mail@example.com"
    # Dòng OCR đã tự đặt dấu cách: khoảng nhiễu 3 px trong từ không phải ranh giới từ
    chars, boxes = _chars("Teaching Assistant course", {2: 3.0, 5: 3.0, 9: 9.0, 14: 3.0, 18: 9.0})
    assert respace(chars, boxes) == "Teaching Assistant course"


def test_fix_terms_only_on_list_lines_and_only_unambiguous():
    vocab = [*static_vocabulary(), "Milvus"]
    assert fix_terms("Tools: Kubemetes, Dockr, Scikit-leam, Mivus", vocab) == (
        "Tools: Kubernetes, Docker, Scikit-Learn, Milvus"
    )
    prose = "We reach our goals and react quickly to feedback."
    assert fix_terms(prose, vocab) == prose


def test_section_heading_with_one_misread_character():
    assert _match_section("HQC VAN") == ("education", True)
    assert _match_section("SKlLLS") == ("skills", True)
    assert _match_section("EXPERlENCE") == ("experience", True)
    assert _match_section("Skillful leader") == (None, False)
    assert _match_section("Proficient") == (None, False)


def _box(x0: float, x1: float, top: float, text: str, h: float = 12.0) -> dict:
    return {"x0": x0, "x1": x1, "top": top, "bottom": top + h, "text": text}


def test_two_column_reading_order_and_name_first():
    sidebar = [
        _box(20, 150, 120 + 20 * i, t) for i, t in enumerate(["CONTACT", "an@example.com", "SKILLS", "Python", "SQL"])
    ]
    main = [
        _box(220, 560, 125 + 20 * i, t)
        for i, t in enumerate(
            [
                "SUMMARY",
                "Data student building dashboards",
                "EXPERIENCE",
                "Beta Retail JSC analyst intern",
                "Built weekly sales dashboard",
            ]
        )
    ]
    name = [_box(220, 420, 40, "Pham Quang Huy", h=26), _box(20, 120, 40, "XXX", h=30)]
    lines = [line for line in extract_layout_aware_text(sidebar + main + name, 600, 800).splitlines() if line.strip()]
    assert lines[0] == "Pham Quang Huy"  # chữ to nhất có dạng họ tên; dải che "XXX" không phải họ tên
    assert (
        lines.index("SKILLS")
        < lines.index("SQL")
        < lines.index("SUMMARY")
        < lines.index("Built weekly sales dashboard")
    )


def test_right_aligned_dates_stay_on_the_same_line():
    rows = [
        (
            _box(40, 330, 100 + 30 * i, f"University of Information Technology {i}"),
            _box(470, 560, 100 + 30 * i, "09/2021 – 06/2025"),
        )
        for i in range(4)
    ]
    lines = extract_layout_aware_text([b for pair in rows for b in pair], 600, 800).splitlines()
    assert "University of Information Technology 0 09/2021 – 06/2025" in lines


_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
_MC = 'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'


def test_docx_reads_text_boxes_once_in_document_order():
    doc = Document()
    doc.add_paragraph("Tran Minh Quan")
    box = "<w:txbxContent><w:p><w:r><w:t>CONTACT</w:t></w:r></w:p><w:p><w:r><w:t>quan@example.com</w:t></w:r></w:p></w:txbxContent>"
    # Word lưu hộp chữ hai bản: DrawingML (mc:Choice) và VML dự phòng (mc:Fallback)
    doc.add_paragraph()._p.append(
        parse_xml(
            f'<w:r {_W} {_MC}><mc:AlternateContent><mc:Choice Requires="wps"><w:pict>{box}</w:pict></mc:Choice>'
            f"<mc:Fallback><w:pict>{box}</w:pict></mc:Fallback></mc:AlternateContent></w:r>"
        )
    )
    doc.add_paragraph("EXPERIENCE")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Tools", "Excel, SQL"
    doc.add_paragraph("EDUCATION")
    buf = io.BytesIO()
    doc.save(buf)
    lines = [line for line in _extract_docx(buf.getvalue()).splitlines() if line.strip()]
    assert lines == ["Tran Minh Quan", "CONTACT", "quan@example.com", "EXPERIENCE", "Tools | Excel, SQL", "EDUCATION"]


def test_garbled_pdf_text_is_detected():
    normal = "Sinh viên năm cuối ngành Marketing, mong muốn trở thành chuyên viên Digital Marketing. " * 3
    assert not looks_garbled(normal)
    assert looks_garbled("(cid:12)(cid:45)(cid:3) " * 40)
    assert looks_garbled("".join(chr(0xE000 + i % 300) for i in range(300)))
    assert readable_chars("ab (cid:12) c") == 3


def test_scan_pages_render_at_a_steady_width():
    assert round(render_scale(595.0) * 595.0) == 1800  # A4
    assert round(render_scale(455.0) * 455.0) == 1800  # ảnh chụp đóng gói khổ nhỏ
    assert render_scale(100.0) == 5.0


def test_no_lone_consonant_pieces_and_split_tool_names_are_fixed():
    # Khoảng trống giả trước "m" (chữ rộng) không tạo mảnh "m anaged"
    chars, boxes = _chars("Resolvedlabsessionsandmanaged", {8: 8.0, 11: 8.0, 19: 8.0, 22: 8.0, 23: 7.5})
    assert respace(chars, boxes) == "Resolved lab sessions and managed"
    vocab = static_vocabulary()
    assert fix_terms("Python,Pandas,Scikit-lea m", vocab) == "Python,Pandas,Scikit-Learn"
    assert fix_terms("Languages: Python, Java, SQL", vocab) == "Languages: Python, Java, SQL"
    lines = ["EDU CATION", "WORK EXPERIENCE", "SKI LLS"]
    assert rejoin_headings("\n".join(lines)).splitlines() == ["EDUCATION", "WORK EXPERIENCE", "SKILLS"]
