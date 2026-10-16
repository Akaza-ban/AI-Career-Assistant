"""Xử lý file CV tải lên: kiểm tra định dạng/dung lượng, trích văn bản, lưu file gốc (US-SV-01)."""

import io
import logging
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from src.api.errors import AppError
from src.config import get_settings
from src.services import ocr

logger = logging.getLogger(__name__)

# Khớp CV_MAX_BYTES ở frontend/components/student/cv/use-cv-upload.ts
CV_MAX_BYTES = 5 * 1024 * 1024
CV_MEDIA_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    # Chỉ để mở lại file ảnh đã tải lên trước đây; tải mới chỉ nhận UPLOAD_TYPES (PDF scan vẫn được OCR).
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
}
UPLOAD_TYPES = ("pdf", "docx")
# Ít hơn ngưỡng này coi như không có lớp văn bản (ví dụ PDF scan ảnh) — US-SV-01 AC2
MIN_TEXT_CHARS = 50
# Lớp chữ ít hơn ngưỡng này thì đọc lại bằng OCR (PDF scan có watermark / vài chữ rác)
OCR_MIN_TEXT_CHARS = 120


@dataclass
class ExtractedCv:
    file_type: str
    text: str
    page_count: int | None
    # True khi chữ được đọc bằng OCR (PDF scan): kết quả phân tích có thể kém chính xác hơn
    read_by_ocr: bool = False


def file_type_of(file_name: str) -> str:
    """Trả về "pdf" | "docx" hoặc báo lỗi 415 nếu sai định dạng (US-SV-01 AC3)."""
    ext = Path(file_name).suffix.lower().lstrip(".")
    if ext not in UPLOAD_TYPES:
        raise AppError(415, "Chỉ hỗ trợ file PDF hoặc DOCX.", "unsupported_file_type")
    return ext


def check_size(size: int) -> None:
    if size > CV_MAX_BYTES:
        raise AppError(413, "File vượt quá dung lượng cho phép (tối đa 5MB).", "file_too_large")
    if size == 0:
        raise AppError(422, "File rỗng. Hãy chọn lại file CV.", "cv_not_extractable")


# Chữ tiếng Việt có dấu thường nằm ở font dự phòng riêng (PDF xuất từ Chrome, Canva...). Khi đó cách trích thông
# thường chèn khoảng trắng quanh từng chữ ("PH Ạ M MINH NH Ậ T"); chế độ "layout" đặt chữ theo toạ độ nên đúng hơn.
_VIET_LETTERS = "ĂăÂâĐđÊêÔôƠơƯư" + "".join(chr(c) for c in range(0x1EA0, 0x1EFA))
_ISOLATED_VIET = re.compile(rf"(?<=\S) [{_VIET_LETTERS}] ?(?=\S)")
_VIET = re.compile(rf"[{_VIET_LETTERS}]")


def _split_diacritics(text: str) -> bool:
    """True khi phần lớn chữ có dấu bị tách rời khỏi chữ bên cạnh bởi khoảng trắng."""
    total = len(_VIET.findall(text))
    return total >= 10 and len(_ISOLATED_VIET.findall(text)) / total > 0.3


def clean_stutter_text(text: str) -> str:
    """Khử hiện tượng lặp đôi ký tự (DDựự áánn -> Dự án, 771155 -> 715) do Canva/PDF export 2 layer chữ đè nhau."""
    if not text:
        return text

    def _clean_line(line: str) -> str:
        line_nfc = unicodedata.normalize("NFC", line)
        tokens = re.split(r"(\s+)", line_nfc)
        has_obvious_stutter = any(
            len(t) >= 4 and len(t) % 2 == 0 and all(t[i] == t[i + 1] for i in range(0, len(t), 2))
            for t in tokens
            if t.strip()
        )
        if not has_obvious_stutter:
            return line_nfc

        def _clean_token(t: str) -> str:
            if not t.strip():
                return t
            n = len(t)
            if n >= 2 and n % 2 == 0 and all(t[i] == t[i + 1] for i in range(0, n, 2)):
                return "".join(t[i] for i in range(0, n, 2))
            return t

        return "".join(_clean_token(t) for t in tokens)

    return "\n".join(_clean_line(line) for line in text.splitlines())


def _extract_pdfplumber_page(page: Any) -> str:
    """Chữ một trang PDF theo thứ tự đọc (chia trang theo khoảng trắng, xem cv_layout.extract_layout_aware_text)."""
    from src.services.cv_layout import extract_layout_aware_text

    words = page.extract_words()
    if not words:
        return (page.extract_text() or "").strip()
    return clean_stutter_text(extract_layout_aware_text(words, page.width, page.height))


def _extract_pdf(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise AppError(422, "File PDF đang được đặt mật khẩu. Hãy gỡ mật khẩu rồi tải lại.", "cv_not_extractable")
    except AppError:
        raise
    except Exception:
        pass

    # 1. Trích xuất bằng pdfplumber kết hợp phân tích bố cục đa cột (layout-aware)
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = pdf.pages
            extracted = [_extract_pdfplumber_page(p) for p in pages]
            plumber_text = "\n\n".join(t for t in extracted if t).strip()
            if len("".join(plumber_text.split())) >= MIN_TEXT_CHARS:
                return plumber_text, len(pages)
    except Exception as e:
        logger.debug("pdfplumber trích xuất không thành công, dùng fallback pypdf: %s", e)

    # 2. Dự phòng pypdf khi pdfplumber không khả dụng hoặc văn bản quá ngắn
    reader = PdfReader(io.BytesIO(data))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    if _split_diacritics(text):
        layout = "\n".join(page.extract_text(extraction_mode="layout") or "" for page in reader.pages)
        text = "\n".join(" ".join(line.split()) for line in layout.splitlines())
    return text, len(reader.pages)


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


def _in_textbox(node: Any) -> bool:
    return any(a.tag == f"{_W}txbxContent" for a in node.iterancestors())


def _paragraph_text(p: Any, *, inside_box: bool = False) -> str:
    """Chữ của một đoạn (w:p); đoạn thường không lấy chữ của hộp chữ con (đọc riêng ở _textboxes)."""
    parts = []
    for node in p.iter(f"{_W}t", f"{_W}tab"):
        if not inside_box and _in_textbox(node):
            continue
        parts.append("\t" if node.tag == f"{_W}tab" else node.text or "")
    return "".join(parts)


def _textboxes(element: Any) -> list[str]:
    """Đoạn trong các hộp chữ nằm trong element. Word lưu mỗi hộp hai bản (DrawingML + VML dự phòng) nên bỏ bản
    nằm trong mc:Fallback để không lặp."""
    lines = []
    for box in element.iter(f"{_W}txbxContent"):
        if any(a.tag == _MC_FALLBACK for a in box.iterancestors()):
            continue
        lines += [_paragraph_text(p, inside_box=True) for p in box.iter(f"{_W}p")]
    return lines


def _extract_docx(data: bytes) -> str:
    """
    Chữ của DOCX theo đúng thứ tự trong file: đoạn văn, bảng, hộp chữ (template hay để liên hệ / kỹ năng trong hộp
    chữ) và phần đầu trang (một số mẫu đặt họ tên, liên hệ ở đó).
    """
    from docx import Document

    document = Document(io.BytesIO(data))
    lines: list[str] = []
    header = document.sections[0].header if document.sections else None
    if header is not None and not header.is_linked_to_previous:
        lines += [_paragraph_text(p._p) for p in header.paragraphs]
        lines += _textboxes(header._element)
    for block in document.element.body.iterchildren():
        if block.tag == f"{_W}p":
            lines.append(_paragraph_text(block))
            lines += _textboxes(block)
        elif block.tag == f"{_W}tbl":
            for row in block.iter(f"{_W}tr"):
                cells = ["\n".join(_paragraph_text(p) for p in cell.iter(f"{_W}p")) for cell in row.iter(f"{_W}tc")]
                lines.append(" | ".join(c for c in cells if c.strip()))
    return "\n".join(lines)


_CID = re.compile(r"\(cid:\d+\)")


def _odd_char(c: str) -> bool:
    """Ký tự không thể là chữ của CV: vùng dành riêng của font, ký tự thay thế, ký tự điều khiển."""
    return c == "�" or unicodedata.category(c) in ("Co", "Cn", "Cc", "Cs")


def readable_chars(text: str) -> int:
    """Số ký tự đọc được (không tính khoảng trắng, mã "(cid:12)" và ký tự lạ)."""
    return sum(1 for c in _CID.sub("", text) if not c.isspace() and not _odd_char(c))


def looks_garbled(text: str) -> bool:
    """Lớp chữ của PDF hỏng bảng mã font: nhiều mã (cid:N) hoặc ký tự lạ, hoặc gần như không có nguyên âm."""
    visible = [c for c in text if not c.isspace()]
    if len(visible) < 120:
        return False
    cid_chars = sum(len(m) for m in _CID.findall(text))
    odd = sum(1 for c in visible if _odd_char(c))
    letters = [c for c in unicodedata.normalize("NFD", text).lower() if c.isalpha()]
    vowels = sum(1 for c in letters if c in "aeiouy")
    return (cid_chars + odd) / len(visible) > 0.2 or (len(letters) > 100 and vowels / len(letters) < 0.15)


def has_text_layer(data: bytes) -> bool:
    """PDF có lớp chữ đọc được (đủ để không cần OCR, cùng ngưỡng với extract_cv)."""
    try:
        text, _ = _extract_pdf(data)
    except Exception:
        return False
    return readable_chars(text) >= OCR_MIN_TEXT_CHARS and not looks_garbled(text)


def extract_cv(file_type: str, data: bytes, vocabulary: list[str] | None = None) -> ExtractedCv:
    """
    Trích văn bản từ CV; tự động kích hoạt OCR khi là PDF scan không có lớp text.
    vocabulary: tên kỹ năng của thư viện JD, dùng sửa tên công cụ OCR đọc nhầm (cùng taxonomy và bảng lộ trình).
    """
    read_by_ocr = False
    try:
        if file_type == "pdf":
            text, page_count = _extract_pdf(data)
            chars_count = readable_chars(text)
            # PDF không có lớp chữ, quá ít chữ (< 120 ký tự: scan có watermark / chữ rác) hoặc lớp chữ hỏng bảng mã
            # font (trông bình thường nhưng trích ra "(cid:12)" / ký tự lạ) thì đọc lại bằng OCR
            if chars_count < OCR_MIN_TEXT_CHARS or looks_garbled(text):
                logger.info("PDF có %d ký tự đọc được. Tiến hành quét OCR qua các trang...", chars_count)
                ocr_text, ocr_pages = ocr.ocr_pdf_bytes(data)
                if readable_chars(ocr_text) > chars_count:
                    text = ocr_text
                    page_count = ocr_pages or page_count
                    read_by_ocr = True
        else:
            text, page_count = _extract_docx(data), None
    except AppError:
        raise
    except Exception as e:
        logger.warning("Không đọc được file CV (%s)", file_type)
        raise AppError(
            422, "File bị lỗi hoặc không đúng định dạng. Hãy kiểm tra lại file.", "cv_not_extractable"
        ) from e

    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    # Chỉ làm sạch cách đọc (chữ lặp đôi do PDF có 2 lớp chữ đè nhau); không để AI viết lại nội dung CV:
    # chữ lưu lại phải là chữ của sinh viên (phân tích và tối ưu đều dựa trên nó).
    text = clean_stutter_text(text)
    if read_by_ocr:
        from src.services.ocr_terms import fix_terms, rejoin_headings, static_vocabulary

        text = fix_terms(rejoin_headings(text), [*static_vocabulary(), *(vocabulary or [])])

    if len("".join(text.split())) < MIN_TEXT_CHARS:
        raise AppError(
            422,
            "Không trích xuất được nội dung CV. Hãy đảm bảo file có chữ rõ ràng hoặc tải lại bản có lớp văn bản.",
            "cv_not_extractable",
        )
    return ExtractedCv(file_type=file_type, text=text, page_count=page_count, read_by_ocr=read_by_ocr)


def merge_wrapped(lines: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Ghép nửa sau của câu bị xuống dòng (cont) vào dòng trước, để mỗi câu là một đoạn."""
    merged: list[dict[str, Any]] = []
    for line in lines:
        if line.get("cont") and merged:
            merged[-1] = {**merged[-1], "text": f"{merged[-1]['text']} {line['text']}"}
        else:
            merged.append(line)
    return merged


def build_docx(document: dict, developing_title: str | None = None) -> bytes:
    """
    File DOCX cho phiên bản CV đã tối ưu, bố cục gọn kiểu Harvard (không giữ thiết kế của file gốc; chữ giữ nguyên):
    họ tên và liên hệ căn giữa, tiêu đề mục in hoa có kẻ dưới, tên mục con in đậm với thời gian căn phải,
    câu mô tả là gạch đầu dòng, cuối cùng là kỹ năng đang phát triển. Loại từng dòng: src/services/cv_layout.py.
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    from src.services.cv_layout import DEVELOPING_TITLES

    doc = Document()
    page = doc.sections[0]
    page.page_width, page.page_height = Inches(8.27), Inches(11.69)  # A4
    page.left_margin = page.right_margin = Inches(0.7)
    page.top_margin = page.bottom_margin = Inches(0.6)
    text_width = page.page_width - page.left_margin - page.right_margin

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.08

    def para(style: str | None = None, *, before: float = 0, align: Any = None):
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(0)
        if align is not None:
            p.alignment = align
        return p

    def run(p: Any, text: str, *, bold: bool = False, italic: bool = False, size: float | None = None):
        r = p.add_run(text)
        r.bold, r.italic = bold or None, italic or None
        if size:
            r.font.size = Pt(size)
        return r

    def heading(title: str) -> None:
        p = para(before=10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = run(p, title, bold=True, size=12)
        r.font.all_caps = True
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
            bottom.set(qn(key), value)
        border.append(bottom)
        p._p.get_or_add_pPr().append(border)

    def labelled(p: Any, text: str) -> None:
        label, _, rest = text.partition(":")
        run(p, f"{label}:", bold=True)
        run(p, rest)

    def line_paragraph(line: dict[str, Any]) -> None:
        kind, text = line.get("kind", "bullet"), line["text"]
        if kind == "entry":
            p = para(before=4)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
            date = line.get("date")
            title = _without_date(text, date) if date else text
            run(p, title, bold=True)
            if date and title != text:
                run(p, f"\t{date}")
        elif kind in ("role", "tech"):
            run(para(), text, italic=True, bold=kind == "tech")
        elif kind == "label":
            labelled(para(), text)
        elif kind == "bullet":
            p = para("List Bullet")
            p.paragraph_format.left_indent = Inches(0.35)
            run(p, text)
        else:
            p = para(align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            run(p, text)

    header = merge_wrapped(document["header"])
    name = next((line for line in header if line.get("kind") == "name"), header[0] if header else None)
    if name:
        run(para(align=WD_ALIGN_PARAGRAPH.CENTER), name["text"], bold=True, size=18).font.all_caps = True
    contacts = [line["text"] for line in header if line is not name and line.get("kind", "contact") == "contact"]
    if contacts:
        run(para(align=WD_ALIGN_PARAGRAPH.CENTER, before=2), " | ".join(contacts), size=10.5)
    for line in header:
        if line is not name and line.get("kind") == "text":
            run(para(align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=6), line["text"])

    for section in document["sections"]:
        heading(section["title"])
        for line in merge_wrapped(section["lines"]):
            line_paragraph(line)
    if document["developing_skills"]:
        heading(developing_title or DEVELOPING_TITLES.get(document.get("lang", "vi"), DEVELOPING_TITLES["vi"]))
        for skill in document["developing_skills"]:
            labelled(para(), f"{skill['name']}: {skill['note']}")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _without_date(text: str, date: str) -> str:
    """Tên của dòng entry: bỏ phần thời gian (ở cuối hoặc đầu dòng) cùng dấu ngăn."""
    at = text.rfind(date)
    if at < 0:
        return text
    title = (text[:at] + text[at + len(date) :]).strip(" \t|,–—-:")
    # "(2022 – 2024)" → bỏ cả cặp ngoặc quanh thời gian
    title = title.removesuffix("(").removeprefix(")").replace("()", "").strip(" \t|,–—-:")
    return title or text


# Ảnh trang PDF để xem CV ngay trong app (không dùng trình xem PDF của trình duyệt): ~150 dpi với khổ A4
PAGE_IMAGE_WIDTH = 1240
MAX_PREVIEW_PAGES = 10


def pdf_page_count(path: Path) -> int:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(path)
    try:
        return len(pdf)
    finally:
        pdf.close()


def render_pdf_page(path: Path, page: int) -> bytes:
    """Ảnh PNG của trang `page` (đếm từ 1), rộng PAGE_IMAGE_WIDTH px."""
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(path)
    try:
        if not 1 <= page <= min(len(pdf), MAX_PREVIEW_PAGES):
            raise AppError(404, "Không có trang này trong CV.", "cv_page_not_found")
        sheet = pdf[page - 1]
        image = sheet.render(scale=PAGE_IMAGE_WIDTH / sheet.get_width()).to_pil()
        buffer = io.BytesIO()
        image.save(buffer, format="PNG", optimize=True)
        return buffer.getvalue()
    finally:
        pdf.close()


def stored_path(owner_id: str, cv_id: str, file_type: str) -> Path:
    return Path(get_settings().upload_dir) / owner_id / f"{cv_id}.{file_type}"


def save_file(owner_id: str, cv_id: str, file_type: str, data: bytes) -> None:
    path = stored_path(owner_id, cv_id, file_type)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)


def delete_file(owner_id: str, cv_id: str, file_type: str) -> None:
    stored_path(owner_id, cv_id, file_type).unlink(missing_ok=True)
